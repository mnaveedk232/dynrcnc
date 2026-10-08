#!/usr/bin/env python3
"""Per-pair predictions of DynRCNC on the 265-pair benchmark, written to an Excel workbook.

Reads the model output (all_predictions.csv, written by dynrcnc.py) and the benchmark file
(benchmark_265pairs.ddg) and writes  benchmark_265pairs_dynrcnc_predictions.xlsx.
The workbook is only written if the predictions reproduce the numbers of the manuscript (use --force to override)."""
import argparse, math, os, sys
import numpy as np, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
ap = argparse.ArgumentParser()
ap.add_argument('--pred', default='~/Desktop/revision_analysis/00_clean_benchmark/test2_out/all_predictions.csv')
ap.add_argument('--ddg', default='~/Desktop/github2/benchmark_265pairs.ddg')
ap.add_argument('--out', default='~/Desktop/github2/benchmark_265pairs_dynrcnc_predictions.xlsx')
ap.add_argument('--force', action='store_true')
a = ap.parse_args(); P, D, O = [os.path.expanduser(x) for x in (a.pred, a.ddg, a.out)]
pred = pd.read_csv(P); ddg = pd.read_csv(D, sep='\t', on_bad_lines='skip'); ddg.columns = ddg.columns.str.strip()
assert len(ddg) == 265, f'{D} has {len(ddg)} rows, expected 265'
ddg = ddg.rename(columns={'Mutation_Double': 'mutation'}); ddg['mutation'] = ddg.mutation.str.replace(' ', '')
pred['mutation'] = pred.mutation.str.replace(' ', '')
key = ['PDB', 'mutation', 'pH', 'Method', 'dddG']
for d in (pred, ddg): d['k'] = d.groupby(key).cumcount()
m = pred.merge(ddg[key + ['k', 'DDG_Double', 'DDG_Single1', 'DDG_Single2']], on=key + ['k'], how='left')
if len(m) != 265 or m.DDG_Double.isna().any(): sys.exit(f'Stop: {len(m)} rows after the join, {int(m.DDG_Double.isna().sum())} without a benchmark match. Check that --pred and --ddg belong together.')
ORDER = ['1STN', '1BNI', '2LZM', '1PGA', '1CSP', '2RN2', '2CI2']
m['o'] = m.PDB.map({p: i for i, p in enumerate(ORDER)}); m = m.sort_values(['o', 'dddG'], ascending=[True, False]).drop(columns=['o', 'k']).reset_index(drop=True)
def stats(df):
    y = (df.exp == 'Non-additive').values; p = (df.pred == 'Non-additive').values
    tp, tn, fp, fn = int((y & p).sum()), int((~y & ~p).sum()), int((~y & p).sum()), int((y & ~p).sum()); den = math.sqrt(float((tp + fp) * (tp + fn) * (tn + fp) * (tn + fn)))
    return dict(N=len(df), NA=int(y.sum()), ADD=int((~y).sum()), TP=tp, TN=tn, FP=fp, FN=fn, Sens=round(100 * tp / (tp + fn), 1) if tp + fn else None, Spec=round(100 * tn / (tn + fp), 1) if tn + fp else None, MCC=round((tp * tn - fp * fn) / den, 3) if den else 0.0)
tot = stats(m); expected = dict(TP=40, TN=184, FP=27, FN=14, MCC=0.568)
ok = all(tot[k] == v for k, v in expected.items())
print('combined:', tot); print('manuscript:', expected, '->', 'MATCH' if ok else 'DIFFERENT')
if not ok and not a.force: sys.exit('Stop: the predictions do not reproduce the manuscript numbers, so the workbook was not written (add --force to write it anyway).')
ARI = lambda **k: Font(name='Arial', size=10, **k); THIN = Side(style='thin', color='BFBFBF'); BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR = PatternFill('solid', fgColor='1F4E78'); RED = PatternFill('solid', fgColor='F8CBAD'); GRN = PatternFill('solid', fgColor='C6E0B4'); ORG = PatternFill('solid', fgColor='FFE699')
wb = Workbook(); ws = wb.active; ws.title = 'Predictions'
cols = ['PDB', 'Mutation pair', 'pH', 'Method', 'ΔΔG double (kcal/mol)', 'ΔΔG single 1 (kcal/mol)', 'ΔΔG single 2 (kcal/mol)', 'Exp. ΔΔΔG (kcal/mol)', 'Exp. class', 'DynRCNC prediction', 'Correct', 'Deciding rule', 'Criterion', '|Z-DCCM|', 'Sequence separation', 'SS site 1', 'SS site 2']
for c, h in enumerate(cols, 1):
    x = ws.cell(1, c, h); x.font = Font(name='Arial', size=10, bold=True, color='FFFFFF'); x.fill = HDR; x.border = BORDER; x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
m['crit'] = m.mechanism.str.extract(r'_(C\d)_')[0].fillna('none')
for r, row in enumerate(m.itertuples(index=False), 2):
    vals = [row.PDB, row.mutation, row.pH, row.Method, row.DDG_Double, row.DDG_Single1, row.DDG_Single2, round(float(row.dddG), 2), row.exp, row.pred, 'Yes' if row.pred == row.exp else 'No', row.mechanism, row.crit, round(abs(float(row.zdccm)), 3), int(row.seq_dist), row.ss_i, row.ss_j]
    for c, v in enumerate(vals, 1):
        x = ws.cell(r, c, v.item() if hasattr(v, 'item') else v); x.font = ARI(); x.border = BORDER
        if c != 2 and c != 12: x.alignment = Alignment(horizontal='center')
        if c in (9, 10): x.fill = RED if v == 'Non-additive' else GRN
        if c == 11: x.fill = GRN if v == 'Yes' else ORG
for i, w in enumerate([7, 16, 6, 9, 13, 13, 13, 13, 13, 15, 9, 30, 10, 10, 11, 8, 8], 1): ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = 'C2'; ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{len(m) + 1}'; ws.row_dimensions[1].height = 42
sm = wb.create_sheet('Summary'); sm.cell(1, 1, 'DynRCNC on the 265-pair benchmark: confusion matrix and metrics (resubstitution)').font = ARI(bold=True)
hd = ['Protein', 'Pairs', 'Non-additive', 'Additive', 'TP', 'TN', 'FP', 'FN', 'Sensitivity (%)', 'Specificity (%)', 'MCC']
for c, h in enumerate(hd, 1):
    x = sm.cell(3, c, h); x.font = Font(name='Arial', size=10, bold=True, color='FFFFFF'); x.fill = HDR; x.border = BORDER; x.alignment = Alignment(horizontal='center', wrap_text=True)
rows = [(p, stats(m[m.PDB == p])) for p in ORDER] + [('All proteins', tot)]
for r, (nm, s) in enumerate(rows, 4):
    for c, v in enumerate([nm, s['N'], s['NA'], s['ADD'], s['TP'], s['TN'], s['FP'], s['FN'], s['Sens'], s['Spec'], s['MCC']], 1):
        x = sm.cell(r, c, v); x.font = ARI(bold=(nm == 'All proteins')); x.border = BORDER; x.alignment = Alignment(horizontal='center')
        if nm == 'All proteins': x.fill = GRN
for i, w in enumerate([16, 8, 13, 10, 7, 7, 7, 7, 15, 15, 9], 1): sm.column_dimensions[get_column_letter(i)].width = w
lp = os.path.join(os.path.dirname(P), 'lopocv.csv')
if os.path.exists(lp):
    l = pd.read_csv(lp); r0 = 6 + len(rows); sm.cell(r0, 1, 'Leave-one-protein-out cross-validation (from lopocv.csv)').font = ARI(bold=True)
    for c, h in enumerate(l.columns, 1):
        x = sm.cell(r0 + 1, c, h); x.font = Font(name='Arial', size=10, bold=True, color='FFFFFF'); x.fill = HDR; x.border = BORDER; x.alignment = Alignment(horizontal='center', wrap_text=True)
    for r, row in enumerate(l.itertuples(index=False), r0 + 2):
        for c, v in enumerate(row, 1):
            x = sm.cell(r, c, round(float(v), 4) if isinstance(v, (float, np.floating)) else (v.item() if hasattr(v, 'item') else v)); x.font = ARI(); x.border = BORDER; x.alignment = Alignment(horizontal='center')
    if {'N', 'MCC'} <= set(l.columns): print('weighted LOPOCV MCC from lopocv.csv: %.3f (manuscript 0.422)' % ((l.MCC * l.N).sum() / l.N.sum()))
rd = wb.create_sheet('README')
for r, (h, t) in enumerate([('Content', 'Per-pair predictions of DynRCNC for the 265 double-mutation pairs of the benchmark (seven proteins).'),
    ('Source', 'Generated by dynrcnc_predictions_265.py from the model output (all_predictions.csv of dynrcnc.py) and benchmark_265pairs.ddg.'),
    ('Experimental class', 'Non-additive when |ΔΔΔG| ≥ 1.0 kcal/mol, where ΔΔΔG = ΔΔG(double) − ΔΔG(single 1) − ΔΔG(single 2); additive otherwise.'),
    ('Deciding rule / criterion', 'The rule that fired (R1 to R14) and the criterion (C1 to C6) it belongs to; "none" means no rule fired and the pair is predicted additive by default.'),
    ('|Z-DCCM|', 'Absolute within-separation Z-score of the dynamic cross-correlation of the two sites, from the wild-type MD trajectory.'),
    ('Repeated measurements', 'A pair measured under several conditions (pH, method) appears once per measurement; its prediction does not depend on the condition.'),
    ('Columns shaded', 'Red: non-additive; green: additive. Correct: green = yes, orange = no.')], 1):
    rd.cell(r, 1, h).font = ARI(bold=True); x = rd.cell(r, 2, t); x.font = ARI(); x.alignment = Alignment(wrap_text=True, vertical='top')
rd.column_dimensions['A'].width = 26; rd.column_dimensions['B'].width = 110
os.makedirs(os.path.dirname(O), exist_ok=True); wb.save(O); print('written:', O, '| 265 rows | sheets:', wb.sheetnames)
