#!/usr/bin/env python3
"""
Build the Rosetta vs Benchmark vs DynRCNC comparison Excel
(4 Rosetta-completed proteins: 1PGA, 1CSP, 2CI2, 1BNI).

INPUT FILES (put this script in the folder with your Rosetta .txt files):
  - {pdb}_s.txt   # singles grep output, e.g. 1csp_s.txt  (lowercase pdb)
  - {pdb}_d.txt   # doubles grep output, e.g. 1csp_d.txt
       for pdb in: 1pga 1csp 2ci2 1bni
  - all_predictions.csv   # DynRCNC output (same file used for the DDGUN sheet)

Each .txt is the raw `grep Aver *.o -A 1` output, lines like:
  s_A46K.o:Average ddG: **-9.60 REU**
  d_A46K_S48R.o:Average ddG: **-13.23 REU**

OUTPUT:
  - Rosetta_vs_Benchmark_vs_DynRCNC_4proteins.xlsx  (3 sheets)

Run:  python3 build_rosetta_comparison.py
Needs: pip install openpyxl
"""

import csv, re, math, os
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ----------------------------------------------------------------------
# EDIT if needed
# ----------------------------------------------------------------------
PROTEINS    = ['1PGA', '1CSP', '2CI2', '1BNI']   # order in the output
DATA_DIR    = "."                                 # folder with the .txt files
DYNRCNC_CSV = "all_predictions.csv"
OUTFILE     = "Rosetta_vs_Benchmark_vs_DynRCNC_4proteins.xlsx"

# ----------------------------------------------------------------------
# Parse a Rosetta grep .txt file -> {mutation_or_pair : ddG_float}
#   singles file keys look like  'A46K'
#   doubles file keys look like  'A46K_S48R'
# ----------------------------------------------------------------------
LINE_RE = re.compile(r'^[sd]_(?P<name>[A-Za-z0-9_]+)\.o:Average ddG:\s*\**\s*(?P<val>-?\d+\.?\d*)')

def parse_rosetta_txt(path):
    out = {}
    with open(path) as f:
        for line in f:
            m = LINE_RE.match(line.strip())
            if m:
                out[m.group('name')] = float(m.group('val'))
    return out

# ----------------------------------------------------------------------
# Load all four proteins' singles + doubles, compute epistasis dddG
#   ros[(PDB, frozenset{m1,m2})] = (m1, s1, m2, s2, double, dddG)
# ----------------------------------------------------------------------
def load_rosetta(data_dir, proteins):
    ros = {}
    singles_all = {}
    for pdb in proteins:
        p = pdb.lower()
        s_path = os.path.join(data_dir, f"{p}_s.txt")
        d_path = os.path.join(data_dir, f"{p}_d.txt")
        singles = parse_rosetta_txt(s_path)
        doubles = parse_rosetta_txt(d_path)
        singles_all[pdb] = singles
        for pair, dv in doubles.items():
            m1, m2 = pair.split('_')
            if m1 not in singles or m2 not in singles:
                print(f"  WARNING {pdb}: missing single for {pair} "
                      f"({m1}={singles.get(m1)}, {m2}={singles.get(m2)})")
                continue
            s1, s2 = singles[m1], singles[m2]
            ros[(pdb, frozenset([m1, m2]))] = (m1, s1, m2, s2, dv, round(dv - s1 - s2, 2))
    return ros, singles_all

# ----------------------------------------------------------------------
# DynRCNC predictions (only the 4 Rosetta proteins)
# ----------------------------------------------------------------------
def load_dynrcnc(path, proteins):
    rows = []
    with open(path) as f:
        for r in csv.DictReader(f):
            if r['PDB'] not in proteins:
                continue
            m = r['mutation'].replace('"', '').replace(' ', '')
            rows.append((r['PDB'], m, r['pH'], r['Method'],
                         round(float(r['dddG']), 2), r['exp'], r['pred'],
                         r['correct'], r['mechanism'], r['seq_dist'],
                         r['ss_i'], r['ss_j'], r['zdccm']))
    return rows

# ----------------------------------------------------------------------
# Styling helpers
# ----------------------------------------------------------------------
ARI      = lambda **k: Font(name='Arial', size=k.pop('size', 10), **k)
THIN     = Side(style='thin', color='BFBFBF')
BORDER   = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER   = Alignment(horizontal='center', vertical='center')
HDR_FILL = PatternFill('solid', fgColor='1F4E78')
HDR_FONT = Font(name='Arial', bold=True, color='FFFFFF', size=10)
GRP = {
    'exp': PatternFill('solid', fgColor='DDEBF7'),
    'ros': PatternFill('solid', fgColor='FFF2CC'),
    'dyn': PatternFill('solid', fgColor='E2EFDA'),
}

# ----------------------------------------------------------------------
# DynRCNC confusion-matrix metrics (positive class = non-additive)
# ----------------------------------------------------------------------
def metrics(sub):
    TP = TN = FP = FN = 0
    for x in sub:
        e = (x['exp'] == 'Non-additive')
        p = (x['dyn_pred'] == 'Non-additive')
        if   e and p:         TP += 1
        elif not e and not p: TN += 1
        elif not e and p:     FP += 1
        else:                 FN += 1
    N = TP + TN + FP + FN
    se = TP / (TP + FN) if TP + FN else 0
    sp = TN / (TN + FP) if TN + FP else 0
    pr = TP / (TP + FP) if TP + FP else 0
    f1 = 2 * pr * se / (pr + se) if pr + se else 0
    den = math.sqrt((TP + FP) * (TP + FN) * (TN + FP) * (TN + FN))
    mcc = (TP * TN - FP * FN) / den if den else 0
    return TP, TN, FP, FN, N, se * 100, sp * 100, f1, mcc

# ----------------------------------------------------------------------
def build(rows, singles_all):
    rows.sort(key=lambda r: (PROTEINS.index(r['PDB']), -r['exp_dddG']))
    wb = Workbook()

    # ---- Sheet 1: per-pair ----------------------------------------------
    ws = wb.active; ws.title = 'Rosetta per-pair (4 proteins)'
    headers = ['PDB', 'Mutation pair', 'pH', 'Method',
               'Exp dddG (kcal/mol)', 'Exp class',
               'Rosetta ddG s1 (REU)', 'Rosetta ddG s2 (REU)',
               'Rosetta ddG double (REU)', 'Rosetta dddG (REU)',
               'DynRCNC pred', 'DynRCNC correct',
               'seq_dist', 'SS_i', 'SS_j', 'zDCCM']
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h); cell.fill = HDR_FILL; cell.font = HDR_FONT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    r = 2
    for x in rows:
        expw = 'Non-additive' if x['exp'] == 'Non-additive' else 'Additive'
        vals = [x['PDB'], x['pair'], x['pH'], x['Method'],
                x['exp_dddG'], expw,
                x['s1'], x['s2'], x['dbl'], x['ros_dddG'],
                x['dyn_pred'], x['dyn_correct'],
                x['sd'], x['ssi'], x['ssj'], x['zd']]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v); cell.font = ARI(); cell.border = BORDER
            if c != 2:
                cell.alignment = CENTER
            if   5 <= c <= 6:  cell.fill = GRP['exp']
            elif 7 <= c <= 10: cell.fill = GRP['ros']
            elif 11 <= c <= 12: cell.fill = GRP['dyn']
        cc = ws.cell(r, 12)
        cc.font = ARI(color='008000' if cc.value in ('True', True) else 'C00000')
        r += 1

    for i, w in enumerate([7,16,6,9,12,12,13,13,15,13,13,13,9,6,6,9], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    # ---- Sheet 2: DynRCNC metrics ---------------------------------------
    by = defaultdict(list)
    for x in rows:
        by[x['PDB']].append(x)

    ws2 = wb.create_sheet('DynRCNC metrics (4 proteins)')
    ws2.cell(1, 1, 'DynRCNC non-additivity detection (Rosetta-completed proteins)').font = ARI(bold=True, size=12)
    for c, h in enumerate(['PDB','N','TP','TN','FP','FN','Sens%','Spec%','F1','MCC'], 1):
        cell = ws2.cell(2, c, h); cell.fill = HDR_FILL; cell.font = HDR_FONT
        cell.border = BORDER; cell.alignment = CENTER
    rr = 3
    for p in PROTEINS:
        d = metrics(by[p])
        vals = [p, d[4], d[0], d[1], d[2], d[3],
                round(d[5],1), round(d[6],1), round(d[7],3), round(d[8],3)]
        for c, v in enumerate(vals, 1):
            cell = ws2.cell(rr, c, v); cell.font = ARI(); cell.border = BORDER
            cell.alignment = CENTER; cell.fill = GRP['dyn']
        rr += 1
    d = metrics(rows)
    vals = ['COMBINED', d[4], d[0], d[1], d[2], d[3],
            round(d[5],1), round(d[6],1), round(d[7],3), round(d[8],3)]
    for c, v in enumerate(vals, 1):
        cell = ws2.cell(rr, c, v); cell.font = ARI(bold=True); cell.alignment = CENTER
        cell.border = Border(top=Side(style='medium'), bottom=Side(style='medium'), left=THIN, right=THIN)
        cell.fill = PatternFill('solid', fgColor='C6E0B4')
    for i, w in enumerate([10,5,5,5,5,5,8,8,7,8], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # ---- Sheet 3: Rosetta singles reference -----------------------------
    ws3 = wb.create_sheet('Rosetta singles')
    for c, h in enumerate(['PDB', 'Mutation', 'Rosetta ddG (REU)'], 1):
        cell = ws3.cell(1, c, h); cell.fill = HDR_FILL; cell.font = HDR_FONT
        cell.border = BORDER; cell.alignment = CENTER
    r3 = 2
    for p in PROTEINS:
        for mut, v in singles_all[p].items():
            for c, val in enumerate([p, mut, v], 1):
                cell = ws3.cell(r3, c, val); cell.font = ARI(); cell.border = BORDER
                cell.alignment = CENTER
            r3 += 1
    for i, w in enumerate([8, 12, 16], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    ws3.freeze_panes = 'A2'

    wb.save(OUTFILE)
    print(f"Saved {OUTFILE}")
    print(f"  per-pair rows : {len(rows)}")
    print(f"  singles       : {r3 - 2}")
    d = metrics(rows)
    print(f"  DynRCNC (4 proteins): N={d[4]} Sens={d[5]:.1f}% Spec={d[6]:.1f}% MCC={d[8]:.3f}")

# ----------------------------------------------------------------------
if __name__ == '__main__':
    ros, singles_all = load_rosetta(DATA_DIR, PROTEINS)
    dyn = load_dynrcnc(DYNRCNC_CSV, PROTEINS)
    print(f"Rosetta double pairs : {len(ros)}")
    print(f"DynRCNC rows (4 prot): {len(dyn)}")

    rows = []
    missing = 0
    for (pdb, m, pH, Method, exp_dddG, exp, pred, correct, mech, sd, ssi, ssj, zd) in dyn:
        key = (pdb, frozenset(m.split(',')))
        if key not in ros:
            missing += 1
            print(f"  no Rosetta match for {pdb} {m}")
            continue
        m1, s1, m2, s2, dbl, ros_dddG = ros[key]
        rows.append({'PDB': pdb, 'pair': m, 'pH': pH, 'Method': Method,
                     'exp_dddG': exp_dddG, 'exp': exp,
                     's1': s1, 's2': s2, 'dbl': dbl, 'ros_dddG': ros_dddG,
                     'dyn_pred': pred, 'dyn_correct': correct,
                     'sd': sd, 'ssi': ssi, 'ssj': ssj, 'zd': zd})
    if missing:
        print(f"WARNING: {missing} DynRCNC rows had no Rosetta match")

    build(rows, singles_all)
